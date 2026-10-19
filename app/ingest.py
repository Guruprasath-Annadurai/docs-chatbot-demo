"""Reading documents and splitting them into chunks."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, asdict
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Section:
    """A piece of a document before chunking. PDFs give one section per page."""

    text: str
    page: int | None = None


@dataclass
class Chunk:
    source: str
    chunk_index: int
    text: str
    page: int | None = None

    def to_dict(self) -> dict:
        return asdict(self)


# ---------------------------------------------------------------- loading


def load_document(filename: str, data: bytes) -> list[Section]:
    """Turn raw file bytes into text sections, based on the file extension."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"'{filename}' is not supported. Use one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )
    if ext == ".pdf":
        return _load_pdf(data)
    if ext == ".docx":
        return _load_docx(data)
    return [Section(text=data.decode("utf-8", errors="replace"))]


def _load_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(Section(text=text, page=number))
    return sections


def _load_docx(data: bytes) -> list[Section]:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        # Keep Word headings recognisable so the chunker can split on them.
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading"):
            text = f"## {text}"
            lines.append("")
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        lines.append("")
    return [Section(text="\n".join(lines))]


# --------------------------------------------------------------- chunking

_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+\S")
_SENTENCE_END = re.compile(r"(?<=[.!?])\s+")


def _blocks(text: str) -> list[str]:
    """Split text into paragraph-like blocks, keeping headings as their own block."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks: list[str] = []
    for raw in re.split(r"\n\s*\n", text):
        current: list[str] = []
        for line in raw.split("\n"):
            if _HEADING.match(line):
                if current:
                    blocks.append("\n".join(current).strip())
                    current = []
                blocks.append(line.strip())
            else:
                current.append(line.rstrip())
        if current and "\n".join(current).strip():
            blocks.append("\n".join(current).strip())
    return [b for b in blocks if b]


def _split_long(block: str, size: int, overlap: int) -> list[str]:
    """Split a block that is longer than `size`, on sentence and then word boundaries."""
    pieces: list[str] = []
    for sentence in _SENTENCE_END.split(block):
        if len(sentence) <= size:
            pieces.append(sentence)
            continue
        words, line = sentence.split(), ""
        for word in words:
            if line and len(line) + 1 + len(word) > size:
                pieces.append(line)
                line = word
            else:
                line = f"{line} {word}".strip()
        if line:
            pieces.append(line)

    parts: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + 1 + len(piece) > size:
            parts.append(current)
            tail = _tail(current, overlap)
            current = f"{tail} {piece}".strip() if tail else piece
            if len(current) > size:  # the overlap would push us over, drop it
                current = piece
        else:
            current = f"{current} {piece}".strip()
    if current:
        parts.append(current)
    return parts


def _tail(text: str, overlap: int) -> str:
    """The last `overlap` characters of text, starting on a word boundary."""
    if overlap <= 0 or len(text) <= overlap:
        return "" if overlap <= 0 else text
    tail = text[-overlap:]
    space = tail.find(" ")
    return tail[space + 1 :] if space != -1 else tail


def chunk_text(text: str, chunk_size: int = 700, chunk_overlap: int = 100) -> list[str]:
    """Pack paragraphs into chunks of at most `chunk_size` characters.

    A markdown heading always starts a new chunk, so a FAQ written as one
    heading per question ends up as one question per chunk. Paragraphs longer
    than the chunk size are split on sentences with `chunk_overlap` characters
    carried over between pieces.
    """
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    chunks: list[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current.strip() and not _is_only_headings(current):
            chunks.append(current.strip())
        current = ""

    for block in _blocks(text):
        if _HEADING.match(block):
            # Keep a heading with the content that follows it.
            if current and not _is_only_headings(current):
                flush()
            current = f"{current}\n{block}".strip() if current else block
            continue
        if len(block) > chunk_size:
            heading = current if _is_only_headings(current) else ""
            if not heading:
                flush()
            for part in _split_long(block, chunk_size - len(heading) - 1 if heading else chunk_size, chunk_overlap):
                chunks.append(f"{heading}\n{part}".strip() if heading else part)
            current = ""
            continue
        joiner = "\n" if _is_only_headings(current) else "\n\n"
        candidate = f"{current}{joiner}{block}" if current else block
        if len(candidate) > chunk_size:
            flush()
            current = block
        else:
            current = candidate
    flush()
    return chunks


def _is_only_headings(text: str) -> bool:
    lines = [line for line in text.split("\n") if line.strip()]
    return bool(lines) and all(_HEADING.match(line) for line in lines)


def make_chunks(source: str, sections: list[Section], chunk_size: int, chunk_overlap: int) -> list[Chunk]:
    chunks: list[Chunk] = []
    for section in sections:
        for text in chunk_text(section.text, chunk_size, chunk_overlap):
            chunks.append(Chunk(source=source, chunk_index=len(chunks), text=text, page=section.page))
    return chunks
