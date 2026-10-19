import io

import docx
import pytest

from app.ingest import UnsupportedFileType, chunk_text, load_document, make_chunks


def test_headings_start_new_chunks():
    text = "# FAQ\n\n### Do you deliver?\nYes, within 8 km.\n\n### Do you take cards?\nYes, all cards."
    chunks = chunk_text(text, chunk_size=500, chunk_overlap=50)
    assert len(chunks) == 2
    assert chunks[0].startswith("# FAQ\n### Do you deliver?\nYes")
    assert "Yes, within 8 km." in chunks[0]
    assert chunks[1].startswith("### Do you take cards?")


def test_small_paragraphs_are_packed_together():
    text = "First paragraph.\n\nSecond paragraph.\n\nThird paragraph."
    assert chunk_text(text, chunk_size=500, chunk_overlap=50) == [
        "First paragraph.\n\nSecond paragraph.\n\nThird paragraph."
    ]


def test_long_text_is_split_within_size_with_overlap():
    sentences = [f"Sentence number {i} talks about bread and cakes." for i in range(60)]
    chunks = chunk_text(" ".join(sentences), chunk_size=300, chunk_overlap=60)
    assert len(chunks) > 5
    assert all(len(c) <= 300 for c in chunks)
    # The start of each chunk repeats the end of the previous one.
    for previous, current in zip(chunks, chunks[1:]):
        assert current.split(" ")[0] in previous


def test_no_text_means_no_chunks():
    assert chunk_text("   \n\n  ", 500, 50) == []
    assert chunk_text("# Just a heading", 500, 50) == []


def test_overlap_must_be_smaller_than_size():
    with pytest.raises(ValueError):
        chunk_text("hello", chunk_size=100, chunk_overlap=100)


def test_pdf_chunks_keep_page_numbers():
    data = (load_document("opening-hours.pdf", open("sample_docs/opening-hours.pdf", "rb").read()))
    chunks = make_chunks("opening-hours.pdf", data, 500, 80)
    assert {c.page for c in chunks} == {1, 2}
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


def test_docx_headings_are_kept():
    document = docx.Document()
    document.add_heading("Prices", level=2)
    document.add_paragraph("Croissant: 110")
    buffer = io.BytesIO()
    document.save(buffer)
    sections = load_document("prices.docx", buffer.getvalue())
    assert "## Prices" in sections[0].text
    assert "Croissant: 110" in sections[0].text


def test_unsupported_type_is_rejected():
    with pytest.raises(UnsupportedFileType):
        load_document("photo.jpg", b"\xff\xd8")
